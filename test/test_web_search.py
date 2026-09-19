"""搜索用例（数据驱动，读 date/search_testdata.xlsx）。

H5 条数断言用“H5列表期望数(DB)”列；“是否空态=是”断言空态；
KW-12（H5否/ES是）走 DB/ES 差异用例。
"""
from __future__ import annotations

import pytest

from date import (
    SEARCH_DATA_XLSX as XLSX_PATH,
    SHEET_SEARCH_HIST as HIST_SHEET,
    SHEET_SEARCH_KW as KW_SHEET,
)
from common.log_decorator import log_class
from pages.web.search_page import SearchPage



def _load_kw_rows():
    import openpyxl

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[KW_SHEET]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        case_id = str(r[0]).strip()
        if not case_id.startswith("KW-"):
            continue
        rows.append(
            {
                "case_id": case_id,
                "keyword": str(r[1] or "").strip(),
                "db_total": int(r[3] or 0),
                "es_total": int(r[4] or 0),
                "is_empty": str(r[5] or "").strip(),  # 否 / 是 / H5否/ES是
                "map_to": str(r[6] or "").strip(),
            }
        )
    return rows


def _load_hist_expect(state_id: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb[HIST_SHEET]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and str(r[0]).strip() == state_id:
            return str(r[2] or "").strip()
    return ""


try:
    _KW_ROWS = _load_kw_rows()
except FileNotFoundError:
    _KW_ROWS = []
    pytestmark = pytest.mark.skip(reason=f"找不到数据文件 {XLSX_PATH}")

LIST_CASES = [c for c in _KW_ROWS if c["map_to"] == "test_search_jumps_to_list"]
EMPTY_CASES = [c for c in _KW_ROWS if c["map_to"] == "test_search_empty"]
DIFF_CASES = [c for c in _KW_ROWS if c["map_to"] == "test_search_db_es_diff"]


@pytest.mark.web
@pytest.mark.e2e
@log_class
class TestWebSearch:
    def test_open_shows_empty(self, pw_page, config):
        """HIST-01：清空历史后打开，空态文案来自“搜索历史状态”表。"""
        expect_text = _load_hist_expect("HIST-01") or "暂无搜索记录"
        page = SearchPage(pw_page, base_url=config["web"]["base_url"])
        page.open()
        page.clear_search_history()
        page.open()  # 清完重进，确保空态
        assert page.is_visible(SearchPage.SEARCH_INPUT)
        assert page.is_visible(SearchPage.EMPTY_SECTION)
        assert page.get_empty_text() == expect_text

    @pytest.mark.parametrize("case", LIST_CASES, ids=lambda c: f'{c["case_id"]}-{c["keyword"]}')
    def test_search_jumps_to_list(self, pw_page, config, case):
        """有结果词：跳列表 + 条数 == H5列表期望数(DB)。"""
        page = SearchPage(pw_page, base_url=config["web"]["base_url"])
        page.open().search(case["keyword"])
        page.wait_for_url("**/product/list*")
        page.wait_for(f"{SearchPage.RESULT_ITEM} >> nth=0")  # 列表异步加载
        assert page.get_result_count() == case["db_total"], (
            f'{case["case_id"]} {case["keyword"]} 期望DB数{case["db_total"]}'
        )

    @pytest.mark.parametrize("case", EMPTY_CASES, ids=lambda c: f'{c["case_id"]}-{c["keyword"]}')
    def test_search_empty(self, pw_page, config, case):
        """无结果词（是否空态=是）：跳列表 + 等渲染 + 空态出现 + 条数为0。"""
        page = SearchPage(pw_page, base_url=config["web"]["base_url"])
        page.open().search(case["keyword"])
        page.wait_for_url("**/product/list*")
        page.sleep(1500)  # 等列表渲染，避免“还没渲染就数到0”的假通过
        assert page.get_result_count() == 0, (
            f'{case["case_id"]} {case["keyword"]} 期望空态0条，实际{page.get_result_count()}条'
        )
        assert page.is_empty_state()

    @pytest.mark.parametrize("case", DIFF_CASES, ids=lambda c: f'{c["case_id"]}-{c["keyword"]}')
    def test_search_db_es_diff(self, pw_page, config, case):
        """KW-12 单字差异：H5(DB模糊)有结果，ES无分词命中。H5只断言DB数。"""
        page = SearchPage(pw_page, base_url=config["web"]["base_url"])
        page.open().search(case["keyword"])
        page.wait_for_url("**/product/list*")
        page.wait_for(f"{SearchPage.RESULT_ITEM} >> nth=0")
        assert page.get_result_count() == case["db_total"]
        # ES差异由接口用例拿 case["es_total"] 去断言，这里只记录备查：
        # 如 KW-12 es_total 应为 0
